#!/usr/bin/env python3
"""
Build an Excel summary workbook that embeds generated plot images.

Inputs:
  - plots/*.png in repo root (produced by plot_from_master.py; legacy Product_Data_File/plots still read)

Output:
  - plots/plots_summary.xlsx (preferred; xlsxwriter)
    Fallback: same path via openpyxl if xlsxwriter unavailable.

Behavior:
  - One sheet per plot image, sheet named from the plot filename (sanitized, <=31 chars).
  - Adds an Index sheet with links to each plot sheet.
  - If no plots present, prints a warning and exits 0.
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path
from typing import List


ROOT = Path(__file__).resolve().parents[2]
EXPORTS_NEW = ROOT
EXPORTS_LEGACY = ROOT / "Product_Data_File"
PLOTS_DIR = EXPORTS_NEW / "plots"
LEGACY_PLOTS_DIR = EXPORTS_LEGACY / "plots"
OUT_XLSX = PLOTS_DIR / "plots_summary.xlsx"
LEGACY_OUT_XLSX = LEGACY_PLOTS_DIR / "plots_summary.xlsx"


def _prefer_new(new_path: Path, legacy_path: Path) -> Path:
    if new_path.exists():
        return new_path
    if legacy_path.exists():
        return legacy_path
    return new_path


def list_plot_images(plots_dir: Path) -> List[Path]:
    if not plots_dir.exists():
        return []
    imgs = [p for p in sorted(plots_dir.iterdir()) if p.suffix.lower() in (".png", ".jpg", ".jpeg")]
    return imgs


def sheet_name_from_file(path: Path, used: set[str]) -> str:
    base = path.stem.strip() or "Plot"
    # Sanitize to Excel sheet name constraints
    name = re.sub(r"[\\/*?:\[\]]+", "_", base)
    # Trim to 31 characters
    name = name[:31] if len(name) > 31 else name
    if not name:
        name = "Plot"
    orig = name
    i = 2
    while name in used:
        suffix = f"_{i}"
        name = (orig[: max(0, 31 - len(suffix))] + suffix)[:31]
        i += 1
    used.add(name)
    return name


def write_with_xlsxwriter(images: List[Path]) -> None:
    import xlsxwriter  # type: ignore

    PLOTS_DIR.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(str(OUT_XLSX))
    try:
        header_fmt = workbook.add_format({"bold": True, "bg_color": "#dde7f5"})
        title_fmt = workbook.add_format({"bold": True, "font_size": 14})
        link_fmt = workbook.add_format({"font_color": "blue", "underline": True})

        ws_index = workbook.add_worksheet("Index")
        ws_index.write(0, 0, "Plot", header_fmt)
        ws_index.write(0, 1, "Sheet", header_fmt)
        ws_index.set_column(0, 0, 45)
        ws_index.set_column(1, 1, 30)

        used: set[str] = set()
        for idx, img in enumerate(images, start=1):
            sheet = sheet_name_from_file(img, used)
            ws = workbook.add_worksheet(sheet)
            ws.write(0, 0, img.stem, title_fmt)
            ws.set_row(0, 24)
            try:
                ws.insert_image(2, 0, str(img), {"object_position": 2})
            except Exception as exc:
                raise RuntimeError(f"Unable to embed image {img}: {exc}") from exc
            ws_index.write_url(idx, 0, f"internal:'{sheet}'!A1", link_fmt, string=img.stem)
            ws_index.write(idx, 1, sheet)
    finally:
        workbook.close()
    print(f"[DONE] Plot summary workbook -> {OUT_XLSX}")


def write_with_openpyxl(images: List[Path]) -> None:
    from openpyxl import Workbook  # type: ignore
    from openpyxl.drawing.image import Image as XLImage  # type: ignore

    PLOTS_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_index = wb.active
    ws_index.title = "Index"
    ws_index.cell(row=1, column=1, value="Plot")
    ws_index.cell(row=1, column=2, value="Sheet")

    used: set[str] = set()
    for idx, img in enumerate(images, start=2):
        sheet = sheet_name_from_file(img, used)
        ws = wb.create_sheet(title=sheet)
        ws.cell(row=1, column=1, value=img.stem)
        try:
            xlimg = XLImage(str(img))
            ws.add_image(xlimg, "A3")
        except Exception as exc:
            raise RuntimeError(
                f"openpyxl was unable to embed {img} (ensure Pillow is installed): {exc}"
            ) from exc
        ws_index.cell(row=idx, column=1, value=img.name)
        ws_index.cell(row=idx, column=2, value=sheet)

    wb.save(OUT_XLSX)
    print(f"[DONE] Plot summary workbook -> {OUT_XLSX}")


def main() -> int:
    plots_root = _prefer_new(PLOTS_DIR, LEGACY_PLOTS_DIR)
    imgs = list_plot_images(plots_root)
    if not imgs:
        print(f"[WARN] No plot images found in {plots_root}. Generate plots first.")
        return 0
    # Ensure new root exists for future writes
    PLOTS_DIR.mkdir(parents=True, exist_ok=True)
    # Try xlsxwriter path
    try:
        write_with_xlsxwriter(imgs)
        return 0
    except Exception as e:
        print(f"[WARN] xlsxwriter path unavailable ({e}); trying openpyxl")
    try:
        write_with_openpyxl(imgs)
        return 0
    except Exception as e:
        print(f"[ERROR] Could not write Excel summary: {e}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
