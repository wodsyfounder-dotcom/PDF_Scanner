#!/usr/bin/env python3
"""
Extract page tables from a PDF into an Excel workbook (one sheet per page).

Grid approach (position -> table):
- Get tokens with positions (PyMuPDF words; EasyOCR fallback per page).
- Normalize to [0..1] coordinates using page width/height.
- Cluster rows via baseline proximity (adaptive y tolerance from median line height).
- Detect column separators from large x-gaps aggregated across all rows (robust stats).
- Form column bands; assign tokens to the unique (row, col) cell; join cell tokens.
- Detect simple section header lines and propagate a Section column.
- Choose a header row as the row with widest column coverage; use those as headers.

Usage:
  python scripts/extract_page_tables.py --pdf path/to/file.pdf --pages "1,3-5" [--ocr] [--dpi 600] [--min-conf 0.4] [--out out.xlsx] [--emit-tokens]

Defaults write to tables/<pdf-stem>_tables.xlsx (repo root; legacy Product_Data_File/tables still read if specified)
"""
from __future__ import annotations

import argparse
from pathlib import Path
import math
import os
import re
import sys
from typing import Dict, List, Tuple, Optional


def parse_page_ranges(s: str) -> List[int]:
    if not s:
        return []
    pages: set[int] = set()
    for part in re.split(r"[;,\s]+", s.strip()):
        if not part:
            continue
        if "-" in part:
            a, b = part.split('-', 1)
            try:
                ai = int(re.sub(r"\D", "", a))
                bi = int(re.sub(r"\D", "", b))
                if ai > bi:
                    ai, bi = bi, ai
                for p in range(ai, bi+1):
                    pages.add(p)
            except Exception:
                continue
        else:
            try:
                pages.add(int(re.sub(r"\D", "", part)))
            except Exception:
                continue
    return sorted(pages)


_AERO_UNITS = (
    "%|ppm|ppb|ms|s|sec|kg|g|mg|ug|lb|lbm|lbf|lbs|"
    "N|kN|mN|Ns|bar|mbar|Pa|kPa|MPa|psi|psia|psig|"
    "mm|cm|m|in|ft|K|degC|degF|C|F"
)
NUMBER_REGEX = re.compile(rf"""
    (?<![A-Za-z0-9_.-])
    [-+]?(?:\d{{1,3}}(?:,\d{{3}})+|\d+)(?:\.\d+)?(?:[eE][+-]?\d+)?  # number
    (?:\s?(?:{_AERO_UNITS}))?                                          # optional unit
    (?![A-Za-z0-9_.-])
""", re.VERBOSE)


def numeric_only(v: Optional[str]) -> Optional[str]:
    if not v:
        return None
    m = re.search(r"[-+]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?:[eE][+-]?\d+)?", v)
    if not m:
        return v
    return m.group(0).replace(',', '')


def extract_units(v: Optional[str]) -> Optional[str]:
    if not v:
        return None
    unit_core = _AERO_UNITS
    m = re.search(r"(?:\s*(" + unit_core + r"))$", v.strip(), flags=re.IGNORECASE)
    return m.group(1) if m else None


def _words_from_pymupdf(pdf: Path, page_num: int) -> Tuple[List[Tuple[float,float,float,float,str,int]], Tuple[float,float]]:
    try:
        import fitz  # PyMuPDF
    except Exception as e:
        raise RuntimeError("PyMuPDF (fitz) is required") from e
    doc = fitz.open(str(pdf))
    try:
        if not (1 <= page_num <= doc.page_count):
            return [], (0.0, 0.0)
        page = doc.load_page(page_num - 1)
        words = page.get_text("words") or []
        out = []
        for w in words:
            try:
                x0, y0, x1, y1, text = float(w[0]), float(w[1]), float(w[2]), float(w[3]), str(w[4])
                ln = int(w[6]) if len(w) >= 7 else int(round(y0))
                if text.strip():
                    out.append((x0, y0, x1, y1, text.strip(), ln))
            except Exception:
                pass
        bbox = page.rect
        return out, (float(bbox.width), float(bbox.height))
    finally:
        try:
            doc.close()
        except Exception:
            pass


def _words_from_easyocr(pdf: Path, page_num: int, dpi: int, langs: List[str], min_conf: float) -> Tuple[List[Tuple[float,float,float,float,str,int]], Tuple[float,float]]:
    try:
        import fitz
    except Exception as e:
        return [], (0.0, 0.0)
    try:
        import easyocr  # type: ignore
    except Exception:
        return [], (0.0, 0.0)
    doc = fitz.open(str(pdf))
    try:
        if not (1 <= page_num <= doc.page_count):
            return [], (0.0, 0.0)
        page = doc.load_page(page_num - 1)
        pix = page.get_pixmap(dpi=max(200, min(800, dpi)))
        from PIL import Image
        import numpy as np  # EasyOCR expects numpy array or path
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        img_np = np.array(img)
        rdr = easyocr.Reader(langs or ['en'], gpu=False, verbose=False)  # type: ignore
        res = rdr.readtext(img_np)  # [(bbox, text, conf), ...]
        out: List[Tuple[float,float,float,float,str,int]] = []
        for it in res:
            try:
                bbox, text, conf = it
                if conf is not None and conf < min_conf:
                    continue
                xs = [float(pt[0]) for pt in bbox]
                ys = [float(pt[1]) for pt in bbox]
                x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
                if str(text).strip():
                    out.append((x0, y0, x1, y1, str(text).strip(), int(round(y0))))
            except Exception:
                pass
        return out, (float(pix.width), float(pix.height))
    finally:
        try:
            doc.close()
        except Exception:
            pass


def _normalize_tokens(words: List[Tuple[float,float,float,float,str,int]], page_w: float, page_h: float) -> List[Dict]:
    out = []
    pw = max(1.0, page_w or 1.0)
    ph = max(1.0, page_h or 1.0)
    for x0, y0, x1, y1, text, _ln in words:
        cx = (x0 + x1) / 2.0
        cy = (y0 + y1) / 2.0
        out.append({
            'x0': x0 / pw, 'y0': y0 / ph, 'x1': x1 / pw, 'y1': y1 / ph,
            'cx': cx / pw, 'cy': cy / ph,
            'w': max(0.0, (x1 - x0) / pw), 'h': max(0.0, (y1 - y0) / ph),
            'text': text,
        })
    return out


def _cluster_rows(tokens: List[Dict], y_tol: Optional[float] = None) -> List[List[Dict]]:
    if not tokens:
        return []
    toks = sorted(tokens, key=lambda t: t['cy'])
    # Adaptive tolerance from median token height
    import statistics as _stats
    hs = [t['h'] for t in toks if t['h'] > 0]
    med_h = _stats.median(hs) if hs else 0.012
    tol = y_tol if (y_tol is not None) else max(0.006, min(0.04, 1.5 * med_h))
    rows: List[List[Dict]] = []
    cur: List[Dict] = []
    last_y = None
    for t in toks:
        if last_y is None or abs(t['cy'] - last_y) <= tol:
            cur.append(t)
            last_y = t['cy'] if last_y is None else (last_y*0.7 + t['cy']*0.3)
        else:
            cur.sort(key=lambda x: x['cx'])
            rows.append(cur)
            cur = [t]
            last_y = t['cy']
    if cur:
        cur.sort(key=lambda x: x['cx'])
        rows.append(cur)
    return rows


def _infer_separators_grid(rows: List[List[Dict]]) -> List[float]:
    """Detect column separators from large gaps across all rows (normalized x in [0..1])."""
    gaps: List[Tuple[float, float]] = []  # (mid, width)
    for ws in rows:
        if len(ws) < 2:
            continue
        for i in range(len(ws) - 1):
            x1 = ws[i]['x1']
            x2 = ws[i+1]['x0']
            width = x2 - x1
            if width <= 0:
                continue
            mid = x1 + width/2.0
            gaps.append((mid, width))
    if not gaps:
        return []
    import statistics as _stats
    widths = [w for _, w in gaps]
    med = _stats.median(widths)
    try:
        q1 = _stats.quantiles(widths, n=4)[0]
        q3 = _stats.quantiles(widths, n=4)[2]
        iqr = max(0.0, q3 - q1)
    except Exception:
        iqr = med
    thresh = max(0.02, med + 0.5 * iqr)  # normalized threshold
    cand = [m for (m, w) in gaps if w >= thresh]
    if not cand:
        return []
    cand.sort()
    # merge nearby midpoints
    merged: List[float] = []
    tol = max(0.01, 0.5 * thresh)
    for x in cand:
        if not merged or abs(x - merged[-1]) > tol:
            merged.append(x)
        else:
            merged[-1] = (merged[-1] + x) / 2.0
    # keep separators that occur in multiple rows
    # approximate frequency by counting row crossings
    seps: List[float] = []
    for s in merged:
        count = 0
        for ws in rows:
            for i in range(len(ws) - 1):
                if ws[i]['x1'] <= s <= ws[i+1]['x0']:
                    count += 1
                    break
        if count >= max(2, int(len(rows) * 0.10)):
            seps.append(s)
    seps.sort()
    return seps


def _assign_columns_grid(rows: List[List[Dict]], seps: List[float]) -> Tuple[int, List[List[List[Dict]]]]:
    bounds = [0.0] + seps + [1.0]
    col_count = max(1, len(bounds)-1)
    table: List[List[List[Dict]]] = []  # rows -> columns -> tokens
    for ws in rows:
        cols = [[] for _ in range(col_count)]
        for t in ws:
            cx = t['cx']
            # find interval
            idx = 0
            while idx < len(bounds)-1 and not (bounds[idx] <= cx < bounds[idx+1]):
                idx += 1
            if idx >= col_count:
                idx = col_count - 1
            cols[idx].append(t)
        table.append(cols)
    return col_count, table


def _choose_header_row(table: List[List[List[Dict]]]) -> Optional[int]:
    best = None
    best_cov = -1
    for ri, cols in enumerate(table):
        cov = sum(1 for c in cols if c)
        if cov > best_cov:
            best_cov = cov
            best = ri
    return best


def _is_section_line(ws: List[Dict], page_w_norm: float) -> bool:
    if not ws:
        return False
    text = " ".join([w['text'] for w in ws])
    span = (ws[-1]['x1'] - ws[0]['x0'])
    if text.endswith(":") and len(text.split()) <= 8:
        return True
    if span / max(1.0, page_w_norm) > 0.7 and len(text.split()) <= 12:
        return True
    return False


def extract_tables_for_pages(pdf: Path, pages: List[int], use_ocr: bool, dpi: int, min_conf: float, langs: List[str], emit_tokens: bool=False) -> Dict[int, Dict]:
    """Build tables per page using block segmentation + header-anchored columns (fallback to global grid)."""
    out: Dict[int, Dict] = {}
    for p in pages:
        words, (pw, ph) = _words_from_pymupdf(pdf, p)
        if not words and use_ocr:
            words, (pw, ph) = _words_from_easyocr(pdf, p, dpi=dpi, langs=langs, min_conf=min_conf)
        tokens = _normalize_tokens(words, pw, ph)
        rows = _cluster_rows(tokens)

        # Segment by section lines
        blocks: List[Tuple[int, int, Optional[str]]] = []
        start = 0
        current_section: Optional[str] = None
        for ri, ws in enumerate(rows):
            if _is_section_line(ws, 1.0):
                if ri > start:
                    blocks.append((start, ri, current_section))
                current_section = " ".join(w['text'] for w in ws).rstrip(':').strip()
                start = ri + 1
        if start < len(rows):
            blocks.append((start, len(rows), current_section))

        page_headers: List[str] = []
        page_rows: List[List[str]] = []

        for a, b, section in blocks:
            sub_rows = rows[a:b]
            if not sub_rows:
                continue
            # Global separators as baseline
            init_seps = _infer_separators_grid(sub_rows)
            col_count, table = _assign_columns_grid(sub_rows, init_seps)
            header_idx = _choose_header_row(table)

            # Improve separators near header row by leveraging dense header tokens
            block_seps = list(init_seps)
            if header_idx is not None:
                hdr = table[header_idx]
                # Use token centers
                centers = [t['cx'] for cell in hdr for t in cell]
                centers.sort()
                # Insert centers that are far from existing separators
                for c in centers:
                    if not block_seps:
                        block_seps.append(c)
                        continue
                    if all(abs(c - s) > 0.03 for s in block_seps):
                        block_seps.append(c)
                        block_seps.sort()
                # Also look at header row gaps
                gaps: List[Tuple[float, float]] = []
                hdr_tokens = sorted([t for cell in hdr for t in cell], key=lambda t: t['cx'])
                for i in range(len(hdr_tokens)-1):
                    x1 = hdr_tokens[i]['x1']
                    x2 = hdr_tokens[i+1]['x0']
                    gaps.append(((x1+x2)/2.0, x2-x1))
                if gaps:
                    import statistics as _stats
                    widths = [w for (_, w) in gaps]
                    med = _stats.median(widths)
                    try:
                        q1 = _stats.quantiles(widths, n=4)[0]
                        q3 = _stats.quantiles(widths, n=4)[2]
                        iqr = max(0.0, q3-q1)
                    except Exception:
                        iqr = med
                    thr = max(0.01, med + 0.5*iqr)
                    mids = [m for (m,w) in gaps if w >= thr]
                    mids.sort()
                    for m in mids:
                        if not block_seps or abs(m - block_seps[-1]) > 0.01:
                            block_seps.append(m)
                        else:
                            block_seps[-1] = (block_seps[-1] + m)/2.0
            if not block_seps:
                block_seps = init_seps

            col_count, table = _assign_columns_grid(sub_rows, block_seps)
            headers = [f"Column_{i+1}" for i in range(col_count)]
            for ri, cols in enumerate(table):
                if header_idx is not None and ri == header_idx:
                    headers = [" ".join(t['text'] for t in sorted(c, key=lambda t: t['cx'])) or f"Column_{i+1}" for i, c in enumerate(cols)]
                    continue
                cells = [" ".join(t['text'] for t in sorted(c, key=lambda t: t['cx'])) for c in cols]
                page_rows.append([section or ""] + cells)

            if not page_headers:
                page_headers = ['Section'] + headers
            else:
                if len(page_headers) - 1 < len(headers):
                    extra = len(headers) - (len(page_headers) - 1)
                    page_headers += [f"Column_{len(page_headers)+i}" for i in range(1, extra+1)]

        if not page_headers:
            page_headers = ['Section', 'Column_1']
        spec = {'columns': page_headers, 'rows': page_rows}
        if emit_tokens:
            spec['tokens'] = [{'text': t['text'], 'x0': t['x0'], 'y0': t['y0'], 'x1': t['x1'], 'y1': t['y1']} for t in tokens]
        out[p] = spec
    return out


def write_excel(tables: Dict[int, Dict], out_path: Path) -> None:
    try:
        import pandas as pd  # type: ignore
    except Exception as e:
        raise RuntimeError("pandas required to write Excel") from e
    try:
        import xlsxwriter  # noqa: F401
        engine = 'xlsxwriter'
    except Exception:
        try:
            import openpyxl  # noqa: F401
            engine = 'openpyxl'
        except Exception as e:
            raise RuntimeError("No Excel writer (xlsxwriter/openpyxl)") from e
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine=engine) as writer:
        for p in sorted(tables.keys()):
            spec = tables[p]
            df = pd.DataFrame(spec['rows'], columns=spec['columns'])
            sheet = f"p{p:03d}"
            df.to_excel(writer, sheet_name=sheet, index=False)
            try:
                ws = writer.sheets[sheet]
                ws.freeze_panes(1, 1)
                for i, col in enumerate(df.columns):
                    try:
                        max_len = int(df[col].astype(str).map(len).max()) if not df.empty else len(col)
                    except Exception:
                        max_len = len(col)
                    if engine == 'xlsxwriter':
                        ws.set_column(i, i, min(60, max(10, max_len + 2)))
            except Exception:
                pass
            # Optional tokens sheet for debugging/audit
            if 'tokens' in spec:
                try:
                    df_tok = pd.DataFrame(spec['tokens'])
                    df_tok.to_excel(writer, sheet_name=f"p{p:03d}_tokens", index=False)
                except Exception:
                    pass


def main() -> None:
    ap = argparse.ArgumentParser(description="Extract tables from pages and write one Excel workbook (sheets per page)")
    ap.add_argument('--pdf', required=True, help='Path to PDF')
    ap.add_argument('--pages', default='', help='Page list, e.g., "1,3-5" (1-indexed). Empty=all pages')
    ap.add_argument('--ocr', action='store_true', help='Force OCR fallback when PyMuPDF words are empty')
    ap.add_argument('--dpi', type=int, default=600, help='OCR DPI when --ocr (default 600)')
    ap.add_argument('--min-conf', type=float, default=0.4, help='Min OCR confidence [0..1] when --ocr')
    ap.add_argument('--langs', default='en', help='OCR languages csv (for EasyOCR), default en')
    ap.add_argument('--out', default='', help='Output .xlsx path (default tables/<stem>_tables.xlsx)')
    ap.add_argument('--emit-tokens', action='store_true', help='Include raw tokens sheets for debugging')
    args = ap.parse_args()

    pdf = Path(args.pdf)
    if not pdf.exists():
        print(f"[ERROR] PDF not found: {pdf}", file=sys.stderr)
        sys.exit(1)
    if args.pages.strip():
        pages = parse_page_ranges(args.pages)
    else:
        try:
            import fitz
            with fitz.open(str(pdf)) as doc:
                pages = list(range(1, int(getattr(doc,'page_count',len(doc)))+1))
        except Exception:
            pages = []
    if not pages:
        print("[ERROR] No pages selected (check --pages)", file=sys.stderr)
        sys.exit(2)

    langs = [s.strip() for s in re.split(r'[;,]', args.langs) if s.strip()]
    tables = extract_tables_for_pages(pdf, pages, use_ocr=bool(args.ocr), dpi=args.dpi, min_conf=args.min_conf, langs=langs, emit_tokens=bool(args.emit_tokens))
    out = Path(args.out) if args.out else (Path('tables') / f"{pdf.stem}_tables.xlsx")
    write_excel(tables, out)
    print(f"[DONE] Wrote page tables -> {out}")


if __name__ == '__main__':
    main()
