#!/usr/bin/env python3
"""
Simple merged-text extraction pipeline.

- Consumes merged OCR artifacts (combined_page.json/combined.txt) produced by pre_ocr_merge.py.
- Uses a lightweight schema with headers:
    Data Group, Term Label, Term, Header, GroupAfter, GroupBefore, Units, Range (min), Range (max)
- Finds the first matching row under an optional Data Group anchor; otherwise uses the first occurrence
  and emits a debug note when multiple matches exist.
- Outputs XLSX (if pandas/openpyxl available) plus CSV fallback under Product_Data_File/run_data_simple/<timestamp>.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple


APP_ROOT = Path(__file__).resolve().parents[1]
CORE_PATH = APP_ROOT / "Application" / "eidp_term_scanner.core.py"
ROOT = APP_ROOT.parent
RUN_BASE = ROOT / "Product_Data_File" / "run_data_simple"


def _load_core():
    spec = importlib.util.spec_from_file_location("eidp_core", CORE_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Unable to load core from {CORE_PATH}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)  # type: ignore[attr-defined]
    return mod


core = _load_core()
normalize = core._normalize_anchor_token
unit_norm = core.normalize_unit_token
derive_pdf_identity = core.derive_pdf_identity
merged_dir_for_pdf = core._merged_output_dir_for_pdf
pre_ocr_and_merge_pdf = core.pre_ocr_and_merge_pdf


@dataclass
class TermSpec:
    data_group: str
    term_label: str
    term: str
    header: str
    group_after: str
    group_before: str
    units: Optional[str]
    range_min: Optional[str]
    range_max: Optional[str]
    report_mode: str = "value"  # "value" (default) or "cell"
    fuzzy_threshold: Optional[float] = None  # optional fuzzy matching threshold (0-1)


def load_terms(path: Path) -> List[TermSpec]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # type: ignore
        except Exception as exc:
            raise RuntimeError("openpyxl is required to read the simple schema Excel") from exc
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [str(c.value or "").strip() for c in ws[1]]
        idx = {h.lower(): i for i, h in enumerate(headers)}
        out: List[TermSpec] = []
        for row in ws.iter_rows(min_row=2):
            vals = [str(c.value or "").strip() for c in row]
            def _get(name: str) -> str:
                i = idx.get(name)
                return vals[i] if i is not None and i < len(vals) else ""
            report = _get("report content") or _get("report") or _get("report mode")
            report = (report or "value").strip().lower()
            if report not in ("value", "cell"):
                report = "value"
            fuzz_raw = _get("fuzzy") or _get("fuzzy match") or ""
            fuzzy_threshold = None
            if fuzz_raw:
                try:
                    fuzzy_threshold = float(fuzz_raw)
                    if not (0.0 < fuzzy_threshold <= 1.0):
                        fuzzy_threshold = None
                except Exception:
                    if fuzz_raw.strip().lower() in ("yes", "y", "true", "on"):
                        fuzzy_threshold = 0.78  # default fuzz when enabled
            out.append(
                TermSpec(
                    data_group=_get("data group"),
                    term_label=_get("term label") or _get("term"),
                    term=_get("term"),
                    header=_get("header"),
                    group_after=_get("groupafter"),
                    group_before=_get("groupbefore"),
                    units=_get("units") or None,
                    range_min=_get("range (min)") or None,
                    range_max=_get("range (max)") or None,
                    report_mode=report,
                    fuzzy_threshold=fuzzy_threshold,
                )
            )
        return [t for t in out if t.term]
    else:
        import csv
        out: List[TermSpec] = []
        with path.open(newline="", encoding="utf-8") as f:
            rdr = csv.DictReader(f)
            for r in rdr:
                report = (r.get("Report Content") or r.get("Report") or r.get("Report Mode") or "").strip().lower() or "value"
                if report not in ("value", "cell"):
                    report = "value"
                fuzz_raw = (r.get("Fuzzy") or r.get("Fuzzy Match") or "").strip()
                fuzzy_threshold = None
                if fuzz_raw:
                    try:
                        fuzzy_threshold = float(fuzz_raw)
                        if not (0.0 < fuzzy_threshold <= 1.0):
                            fuzzy_threshold = None
                    except Exception:
                        if fuzz_raw.strip().lower() in ("yes", "y", "true", "on"):
                            fuzzy_threshold = 0.78
                out.append(
                    TermSpec(
                        data_group=(r.get("Data Group") or "").strip(),
                        term_label=(r.get("Term Label") or r.get("Term") or "").strip(),
                        term=(r.get("Term") or "").strip(),
                        header=(r.get("Header") or "").strip(),
                        group_after=(r.get("GroupAfter") or "").strip(),
                        group_before=(r.get("GroupBefore") or "").strip(),
                        units=(r.get("Units") or "").strip() or None,
                        range_min=(r.get("Range (min)") or "").strip() or None,
                        range_max=(r.get("Range (max)") or "").strip() or None,
                        report_mode=report,
                        fuzzy_threshold=fuzzy_threshold,
                    )
                )
        return [t for t in out if t.term]


def _normalize_tokens(text: str) -> List[str]:
    return [normalize(tok) for tok in text.split() if normalize(tok)]


def _row_has_term(row_tokens: List[str], term: str, fuzzy_threshold: Optional[float] = None, row_text: str = "") -> bool:
    from difflib import SequenceMatcher
    t = normalize(term)
    if not t:
        return False
    if t in row_tokens or any(t in tok for tok in row_tokens):
        return True
    if fuzzy_threshold:
        # Check fuzzy against tokens and full row text
        for tok in row_tokens:
            if SequenceMatcher(None, t, tok).ratio() >= fuzzy_threshold:
                return True
        if row_text:
            if SequenceMatcher(None, t, normalize(row_text)).ratio() >= fuzzy_threshold:
                return True
    return False


def _parse_range(cells: List[str]) -> Tuple[Optional[str], Optional[str]]:
    import re
    rng_min = rng_max = None
    def _as_float(txt: str) -> Optional[float]:
        try:
            return float(txt.replace(",", ""))
        except Exception:
            return None
    for c in cells:
        txt = c or ""
        pm = re.search(r"([-+]?\d[\d,]*(?:\.\d+)?)[ ]*(?:[±\u00b1\u2213]|\+/?-|\+-|\\+-|Â±)[ ]*([-+]?\d[\d,]*(?:\.\d+)?)", txt)
        if pm:
            base_txt, delta_txt = pm.group(1), pm.group(2)
            base = _as_float(base_txt)
            delta = _as_float(delta_txt)
            if base is not None and delta is not None:
                low = base - delta
                high = base + delta
                rng_min = rng_min or str(low)
                rng_max = rng_max or str(high)
            else:
                rng_min = rng_min or base_txt
                rng_max = rng_max or delta_txt
        if rng_min and rng_max:
            break
        m = re.search(r"([<>]=?)\s*([0-9][0-9.,eE+-]*)", txt)
        if m:
            op, val = m.group(1), m.group(2)
            if op.startswith("<"):
                rng_max = rng_max or val
            elif op.startswith(">"):
                rng_min = rng_min or val
        m2 = re.search(r"([0-9][0-9.,eE+-]*)\s*[-–]\s*([0-9][0-9.,eE+-]*)", txt)
        if m2:
            rng_min = rng_min or m2.group(1)
            rng_max = rng_max or m2.group(2)
    return rng_min, rng_max


def _detect_units(cells: List[str]) -> Optional[str]:
    for c in cells:
        for token in c.split():
            u = unit_norm(token)
            if u:
                return u
    return None


def _extract_numeric_candidate(text: str, range_min: Optional[str], range_max: Optional[str]) -> Optional[str]:
    import re
    if not text:
        return None
    nums = re.findall(r"[-+]?\d[\d,]*(?:\.\d+)?(?:[eE][+-]?\d+)?", text)
    def _clean(n: str) -> str:
        return n.replace(",", "")
    rng = {_clean(r) for r in (range_min, range_max) if r}
    if nums:
        for n in nums:
            c = _clean(n)
            if c not in rng:
                return n
        return nums[0]
    return None


def _first_numeric(text: str) -> Optional[float]:
    import re
    if not text:
        return None
    nums = re.findall(r"[-+]?\d[\d,]*(?:\.\d+)?(?:[eE][+-]?\d+)?", text)
    if not nums:
        return None
    try:
        return float(nums[0].replace(",", ""))
    except Exception:
        return None


def _mark_out_of_range(value: Optional[str], rng_min: Optional[str], rng_max: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    """Return (value_with_flag, debug_note) when value is outside [min,max]."""
    if value is None:
        return value, None
    v_num = _first_numeric(value)
    if v_num is None:
        return value, None
    try:
        min_num = float(str(rng_min).replace(",", "")) if rng_min not in (None, "") else None
    except Exception:
        min_num = None
    try:
        max_num = float(str(rng_max).replace(",", "")) if rng_max not in (None, "") else None
    except Exception:
        max_num = None
    out_of_range = False
    if min_num is not None and v_num < min_num:
        out_of_range = True
    if max_num is not None and v_num > max_num:
        out_of_range = True
    if out_of_range:
        flagged = f"{value} (out of range)"
        return flagged, "out of range"
    return value, None


def _derive_units_and_range(headers: List[str], cells: List[str], schema_units: Optional[str], rng_min: Optional[str], rng_max: Optional[str]) -> Tuple[Optional[str], Optional[str], Optional[str], str]:
    """Infer units/range from header-aligned cells and row content; schema overrides, conflicts noted."""
    debug_note = ""
    auto_units = None
    auto_min = None
    auto_max = None
    hdr_norms = [normalize(h or "") for h in headers] if headers else []

    # Header-aligned units/min/max/range columns
    for idx, hn in enumerate(hdr_norms):
        if idx >= len(cells):
            continue
        txt = cells[idx]
        if hn == "units" and not auto_units:
            auto_units = unit_norm(txt) or (txt.strip() if txt.strip() else None)
        if hn in ("range", "tol", "tolerance", "requirement", "req", "spec", "specification", "limit", "limits"):
            lo, hi = _parse_range([txt])
            auto_min = auto_min or lo
            auto_max = auto_max or hi
        if hn in ("min", "minimum", "lower"):
            if txt.strip():
                auto_min = auto_min or txt.strip()
        if hn in ("max", "maximum", "upper"):
            if txt.strip():
                auto_max = auto_max or txt.strip()

    # Row-level units/range fallback
    if not auto_units:
        auto_units = _detect_units(cells)
    if not auto_min and not auto_max:
        lo, hi = _parse_range(cells)
        auto_min = lo or auto_min
        auto_max = hi or auto_max

    # Apply overrides with conflict notes
    final_units = schema_units or auto_units
    if schema_units and auto_units and normalize(schema_units) != normalize(auto_units):
        debug_note += f"conflicting units (schema={schema_units}, auto={auto_units}); "
    final_min = rng_min or auto_min
    final_max = rng_max or auto_max
    if (rng_min or rng_max) and (auto_min or auto_max):
        if rng_min and auto_min and normalize(rng_min) != normalize(auto_min):
            debug_note += f"conflicting range min (schema={rng_min}, auto={auto_min}); "
        if rng_max and auto_max and normalize(rng_max) != normalize(auto_max):
            debug_note += f"conflicting range max (schema={rng_max}, auto={auto_max}); "
    return final_units, final_min, final_max, debug_note.strip()


def _best_value(cells: List[str], headers: List[str], header_hint: str) -> Tuple[Optional[str], Optional[int]]:
    norm_hdrs = [normalize(str(h or "")) for h in headers] if headers else []
    synonyms = ("measuredvalue", "value", "measurement", "result", "reading", "actual", "observed")

    def _find_header_index(targets: List[str]) -> Optional[int]:
        if not norm_hdrs:
            return None
        for t in targets:
            if t in norm_hdrs:
                return norm_hdrs.index(t)
        # allow contains match (e.g., target='value' matching 'measuredvalue')
        for t in targets:
            for i, h in enumerate(norm_hdrs):
                if t and t in h:
                    return i
        return None

    def _pick_by_norm(targets: List[str]) -> Tuple[Optional[str], Optional[int]]:
        idx = _find_header_index(targets)
        if idx is not None and idx < len(cells):
            return cells[idx].strip(), idx
        return None, None

    # 1) Explicit header hint
    if headers and header_hint:
        target = normalize(header_hint)
        if target:
            # Direct match
            val, idx = _pick_by_norm([target])
            if val:
                return val, idx
            # If hint is generic "value", broaden to synonyms
            if target in ("value", "val"):
                val, idx = _pick_by_norm(list(synonyms))
                if val:
                    return val, idx

    # 2) No header hint (or not found): try value-like headers
    if headers:
        val, idx = _pick_by_norm(list(synonyms))
        if val:
            return val, idx

    # No suitable header match found
    return None, None


def _load_flow_pages(combined_page_json: Path) -> List[Dict[str, object]]:
    obj = json.loads(combined_page_json.read_text(encoding="utf-8"))
    pages = obj.get("pages") if isinstance(obj, dict) else None
    if isinstance(pages, list):
        return pages
    return []


def _text_fuzzy_match(text: str, needle: str, threshold: Optional[float]) -> bool:
    if not threshold:
        return normalize(needle) in normalize(text)
    from difflib import SequenceMatcher
    return SequenceMatcher(None, normalize(text), normalize(needle)).ratio() >= threshold


def extract_from_pdf(pdf_path: Path, terms: List[TermSpec]) -> List[Dict[str, object]]:
    pdf_path = Path(pdf_path)
    merged_dir = merged_dir_for_pdf(pdf_path, None)
    manifest_path = merged_dir / "manifest.json"
    combined_page_json = merged_dir / "combined_page.json"
    if not manifest_path.exists() or not combined_page_json.exists():
        pre_ocr_and_merge_pdf(pdf_path, out_dir=merged_dir)
    pages = _load_flow_pages(combined_page_json)
    flow_with_page: List[Tuple[int, Dict[str, object]]] = []
    for idx, p in enumerate(pages, start=1):
        fl = p.get("flow") if isinstance(p, dict) else None
        if isinstance(fl, list):
            for item in fl:
                if isinstance(item, dict):
                    flow_with_page.append((idx, item))
    program, vehicle, serial = derive_pdf_identity(pdf_path)
    results: List[Dict[str, object]] = []

    for spec in terms:
        term_norm = normalize(spec.term)
        header_norm = normalize(spec.header)
        group_norm = normalize(spec.data_group)
        group_after_norm = normalize(spec.group_after)
        group_before_norm = normalize(spec.group_before)
        fuzz = spec.fuzzy_threshold
        start_idx = 0
        if group_norm:
            for i, (_p, item) in enumerate(flow_with_page):
                txt = str(item.get("text") or "")
                if _text_fuzzy_match(txt, group_norm, fuzz):
                    start_idx = i
                    break

        candidates: List[Tuple[int, Dict[str, object], Dict[str, object]]] = []
        for i, (page_num, item) in enumerate(flow_with_page[start_idx:], start=start_idx):
            if group_after_norm:
                txt = str(item.get("text") or "")
                if _text_fuzzy_match(txt, group_after_norm, fuzz):
                    continue
            if group_before_norm:
                txt = str(item.get("text") or "")
                if _text_fuzzy_match(txt, group_before_norm, fuzz):
                    break
            if str(item.get("type") or "") != "table":
                continue
            tb = item.get("table")
            if not isinstance(tb, dict):
                continue
            headers = tb.get("header_cells") if isinstance(tb.get("header_cells"), list) else []
            norm_headers = [normalize(h or "") for h in headers]
            if header_norm:
                hdr_match = any(header_norm == h or (header_norm and header_norm in h) or (h and h in header_norm) for h in norm_headers)
                if not hdr_match and fuzz:
                    from difflib import SequenceMatcher
                    hdr_match = any(SequenceMatcher(None, header_norm, h).ratio() >= fuzz for h in norm_headers)
                if not hdr_match:
                    continue
            rows_raw = tb.get("rows") if isinstance(tb.get("rows"), list) else []
            for r in rows_raw:
                if not isinstance(r, dict):
                    continue
                cells = [str(x or "") for x in (r.get("cells_text") if isinstance(r.get("cells_text"), list) else [])]
                row_text = " ".join(cells)
                row_tokens = _normalize_tokens(row_text)
                if _row_has_term(row_tokens, spec.term, fuzzy_threshold=fuzz, row_text=row_text):
                    candidates.append((page_num, item, {"cells": cells, "headers": headers}))
        debug_note = ""
        if len(candidates) > 1:
            debug_note = "multiple terms in the file found, used the first"
        if not candidates and flow_with_page:
            # Fallback: scan any text block for the term, pick first
            for page_num, item in flow_with_page[start_idx:]:
                txt = str(item.get("text") or "")
                if term_norm and _text_fuzzy_match(txt, term_norm, fuzz):
                    candidates.append((page_num, item, {"cells": [txt], "headers": []}))
                    debug_note = "matched text flow (no table); used first occurrence"
                    break
        if not candidates:
            results.append({
                "pdf_file": pdf_path.name,
                "program_name": program,
                "vehicle_number": vehicle,
                "serial_component": serial,
                "term_label": spec.term_label or spec.term,
                "term": spec.term,
                "header": spec.header,
                "data_group": spec.data_group,
                "report_mode": spec.report_mode,
                "found": False,
                "value": None,
                "units": None,
                "range_min": spec.range_min,
                "range_max": spec.range_max,
                "page": None,
                "debug": "not found",
            })
            continue
        page_num, _item, payload = candidates[0]
        cells = payload["cells"]
        headers = payload["headers"]
        val, col_idx = _best_value(cells, headers, spec.header)
        rng_min = spec.range_min
        rng_max = spec.range_max
        units, rng_min, rng_max, dbg_units = _derive_units_and_range(headers, cells, spec.units, rng_min, rng_max)
        if dbg_units:
            debug_note = (debug_note + " | " if debug_note else "") + dbg_units
        if group_norm:
            debug_note = (debug_note + " | " if debug_note else "") + "matched after data group"
        # Apply report mode
        if spec.report_mode == "value":
            numeric = _extract_numeric_candidate(val, rng_min, rng_max)
            if numeric:
                val = numeric
        else:
            debug_note = (debug_note + " | " if debug_note else "") + "reported full cell"
        # Flag out-of-range values
        val, dbg_range = _mark_out_of_range(val, rng_min, rng_max)
        if dbg_range:
            debug_note = (debug_note + " | " if debug_note else "") + dbg_range
        results.append({
            "pdf_file": pdf_path.name,
            "program_name": program,
            "vehicle_number": vehicle,
            "serial_component": serial,
            "term_label": spec.term_label or spec.term,
            "term": spec.term,
            "header": spec.header,
            "data_group": spec.data_group,
            "report_mode": spec.report_mode,
            "found": bool(val),
            "value": val,
            "units": units,
            "range_min": rng_min,
            "range_max": rng_max,
            "page": page_num,
            "debug": debug_note,
        })
    return results


def write_outputs(rows: List[Dict[str, object]], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / "scan_results_simple.csv"
    fields = [
        "pdf_file", "program_name", "vehicle_number", "serial_component",
        "term_label", "term", "header", "data_group", "report_mode",
        "value", "units", "range_min", "range_max", "page", "debug",
    ]
    import csv
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fields})
    try:
        import pandas as pd  # type: ignore
        df = pd.DataFrame(rows)[fields]
        xlsx_path = out_dir / "scan_results_simple.xlsx"
        df.to_excel(xlsx_path, index=False)
    except Exception:
        pass


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Simple merged-text extraction pipeline.")
    parser.add_argument("--pdf", action="append", dest="pdfs", required=True, help="PDF path (can repeat)")
    parser.add_argument("--terms", required=True, help="Path to simple schema (xlsx/csv)")
    args = parser.parse_args(argv)
    pdfs = [Path(p).expanduser() for p in args.pdfs]
    terms_path = Path(args.terms).expanduser()
    terms = load_terms(terms_path)
    if not terms:
        print("[ERROR] No terms loaded from schema", file=sys.stderr)
        return 2
    from datetime import datetime
    out_dir = RUN_BASE / datetime.now().strftime("%Y%m%d_%H%M%S")
    rows_all: List[Dict[str, object]] = []
    for pdf in pdfs:
        try:
            rows_all.extend(extract_from_pdf(pdf, terms))
        except Exception as exc:
            rows_all.append({
                "pdf_file": pdf.name,
                "program_name": "",
                "vehicle_number": "",
                "serial_component": "",
                "term_label": "",
                "term": "",
                "header": "",
                "data_group": "",
                "found": False,
                "value": None,
                "units": None,
                "range_min": None,
                "range_max": None,
                "page": None,
                "debug": f"error: {exc}",
            })
    write_outputs(rows_all, out_dir)
    print(f"[DONE] Simple extraction written -> {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
