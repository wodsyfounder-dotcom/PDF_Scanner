#!/usr/bin/env python3
"""
Reconstruct a table from EasyOCR tokens JSON by clustering rows and columns
using spatial heuristics, then export to CSV/XLSX and an SVG visualization.

Inputs: a single page tokens JSON produced by easyocr_dump_spatial.py
Outputs beside an out-prefix:
 - <out-prefix>.table.csv (always)
 - <out-prefix>.table.xlsx (if openpyxl is available)
 - <out-prefix>.table.svg (visual grid + cell text)

Usage:
  py -3.11 scripts/build_table_from_tokens.py \
      --tokens-json Product_Data_File/easy_spatial/sn_4444/sn\ 4444_page_1.json \
      --out-prefix Product_Data_File/easy_spatial/sn_4444/sn_4444_page_1 \
      --row-factor 0.6 --gap-mult 2.5 --center-thresh-mult 2.0 --use-first-row-as-header
"""
from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path
from typing import Dict, List, Tuple


Token = Dict[str, float | str]


def load_tokens(path: Path) -> List[Token]:
    data = json.loads(path.read_text(encoding="utf-8"))
    # ensure shape is list[dict]
    if not isinstance(data, list):
        raise SystemExit(f"[ERROR] Expected list of tokens in {path}")
    return data


def compute_stats(tokens: List[Token]) -> Tuple[float, float, float, float]:
    xs0 = [float(t["x0"]) for t in tokens]
    xs1 = [float(t["x1"]) for t in tokens]
    ws = [float(t.get("w", max(0.0, float(t["x1"]) - float(t["x0"])))) for t in tokens]
    hs = [float(t.get("h", max(0.0, float(t["y1"]) - float(t["y0"])))) for t in tokens]
    min_x = min(xs0) if xs0 else 0.0
    max_x = max(xs1) if xs1 else 0.0
    med_w = median(ws) if ws else 10.0
    med_h = median(hs) if hs else 12.0
    return min_x, max_x, med_w, med_h


def median(values: List[float]) -> float:
    vals = sorted(values)
    n = len(vals)
    if n == 0:
        return 0.0
    mid = n // 2
    if n % 2 == 1:
        return float(vals[mid])
    return float((vals[mid - 1] + vals[mid]) / 2.0)


def sort_tokens(tokens: List[Token]) -> List[Token]:
    return sorted(tokens, key=lambda t: (float(t["cy"]), float(t["cx"])))


def cluster_rows(tokens: List[Token], row_factor: float) -> List[List[Token]]:
    if not tokens:
        return []
    tokens_sorted = sort_tokens(tokens)
    med_h = median([float(t.get("h", 0.0)) for t in tokens if float(t.get("h", 0.0)) > 0.0]) or 12.0
    thresh = max(2.0, row_factor * med_h)
    rows: List[List[Token]] = []
    current: List[Token] = []
    current_cy: float | None = None
    for t in tokens_sorted:
        cy = float(t["cy"])
        if current_cy is None:
            current = [t]
            current_cy = cy
        else:
            if abs(cy - current_cy) <= thresh:
                current.append(t)
                current_cy = (current_cy * (len(current) - 1) + cy) / len(current)
            else:
                rows.append(sorted(current, key=lambda x: float(x["cx"])) )
                current = [t]
                current_cy = cy
    if current:
        rows.append(sorted(current, key=lambda x: float(x["cx"])) )
    return rows


def row_segments(row_tokens: List[Token], gap_mult: float, med_w: float) -> List[Tuple[float, List[Token]]]:
    # Group adjacent tokens into horizontal segments by gap threshold
    if not row_tokens:
        return []
    segs: List[List[Token]] = []
    cur: List[Token] = [row_tokens[0]]
    for prev, t in zip(row_tokens, row_tokens[1:]):
        gap = float(t["x0"]) - float(prev["x1"])
        if gap > (gap_mult * med_w):
            segs.append(cur)
            cur = [t]
        else:
            cur.append(t)
    segs.append(cur)
    # segment center is mean cx
    out: List[Tuple[float, List[Token]]] = []
    for seg in segs:
        center = sum(float(s["cx"]) for s in seg) / max(1, len(seg))
        out.append((center, seg))
    return out


def cluster_columns(centers: List[float], center_thresh: float) -> List[float]:
    if not centers:
        return []
    centers_sorted = sorted(centers)
    clusters: List[List[float]] = [[centers_sorted[0]]]
    for c in centers_sorted[1:]:
        if abs(c - clusters[-1][-1]) <= center_thresh:
            clusters[-1].append(c)
        else:
            clusters.append([c])
    # Return mean center per cluster
    return [sum(g) / len(g) for g in clusters]


def build_grid(rows: List[List[Token]], med_w: float, min_x: float, max_x: float,
               gap_mult: float, center_thresh_mult: float) -> Tuple[List[str], List[List[str]]]:
    # 1) derive candidate column centers from all row segments
    all_seg_centers: List[float] = []
    for row in rows:
        segs = row_segments(row, gap_mult=gap_mult, med_w=med_w)
        for c, _ in segs:
            all_seg_centers.append(c)
    if not all_seg_centers:
        return [], []
    center_thresh = max(2.0, center_thresh_mult * med_w)
    col_centers = cluster_columns(all_seg_centers, center_thresh=center_thresh)
    col_centers.sort()
    # 2) Assign row segments to nearest columns
    table_rows: List[List[str]] = []
    for row in rows:
        segs = row_segments(row, gap_mult=gap_mult, med_w=med_w)
        cells = [""] * len(col_centers)
        for c, seg in segs:
            # nearest column
            nearest_idx = min(range(len(col_centers)), key=lambda i: abs(col_centers[i] - c))
            if abs(col_centers[nearest_idx] - c) <= center_thresh:
                text = " ".join(s["text"] for s in seg if str(s.get("text", "")).strip())
                if cells[nearest_idx]:
                    cells[nearest_idx] += " " + text
                else:
                    cells[nearest_idx] = text
        table_rows.append(cells)
    # Header defaults
    headers = [f"col_{i+1}" for i in range(len(col_centers))]
    return headers, table_rows


def write_csv(out_prefix: Path, headers: List[str], rows: List[List[str]]) -> Path:
    csv_path = out_prefix.with_suffix(".table.csv")
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        if headers:
            w.writerow(headers)
        for r in rows:
            w.writerow(r)
    return csv_path


def write_xlsx(out_prefix: Path, headers: List[str], rows: List[List[str]]) -> Path | None:
    try:
        from openpyxl import Workbook  # type: ignore
        from openpyxl.utils import get_column_letter  # type: ignore
    except Exception:
        return None
    xlsx_path = out_prefix.with_suffix(".table.xlsx")
    wb = Workbook()
    ws = wb.active
    if headers:
        ws.append(headers)
    for r in rows:
        ws.append(r)
    # Autosize columns roughly
    max_cols = max((len(headers), *(len(r) for r in rows)), default=0)
    for c in range(1, max_cols + 1):
        col_letter = get_column_letter(c)
        max_len = 8
        for cell in ws[col_letter]:
            max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[col_letter].width = min(60, max_len + 2)
    wb.save(xlsx_path)
    return xlsx_path


def write_svg(out_prefix: Path, headers: List[str], rows: List[List[str]],
              tokens: List[Token], row_factor: float, gap_mult: float,
              center_thresh_mult: float, scale: float = 0.25) -> Path:
    # Recompute the grid geometry to draw cell rectangles consistently
    min_x, max_x, med_w, med_h = compute_stats(tokens)
    clustered_rows = cluster_rows(tokens, row_factor=row_factor)
    # Compute row spans (y0,y1) per clustered row
    row_spans: List[Tuple[float, float]] = []
    for row in clustered_rows:
        y0 = min(float(t["y0"]) for t in row)
        y1 = max(float(t["y1"]) for t in row)
        row_spans.append((y0, y1))
    # Column centers and boundaries
    _, _rows = build_grid(clustered_rows, med_w, min_x, max_x, gap_mult, center_thresh_mult)
    # Derive centers again
    all_seg_centers: List[float] = []
    for row in clustered_rows:
        for c, _ in row_segments(row, gap_mult=gap_mult, med_w=med_w):
            all_seg_centers.append(c)
    center_thresh = max(2.0, center_thresh_mult * med_w)
    col_centers = cluster_columns(all_seg_centers, center_thresh=center_thresh)
    col_centers.sort()
    if not row_spans or not col_centers:
        svg_path = out_prefix.with_suffix(".table.svg")
        out = ["<svg xmlns='http://www.w3.org/2000/svg' width='800' height='200'>",
               "<text x='10' y='20' font-family='monospace' font-size='14'>No grid to draw</text>",
               "</svg>"]
        svg_path.write_text("\n".join(out), encoding="utf-8")
        return svg_path
    # Boundaries as midpoints between centers, and edges with min_x/max_x
    bounds: List[float] = []
    bounds.append((min_x + col_centers[0]) / 2.0)
    for a, b in zip(col_centers, col_centers[1:]):
        bounds.append((a + b) / 2.0)
    bounds.append((col_centers[-1] + max_x) / 2.0)
    # Canvas
    margin = 20.0
    W = (max_x - min_x) * scale + 2 * margin
    H = (max(float(t["y1"]) for t in tokens) - min(float(t["y0"]) for t in tokens)) * scale + 2 * margin
    # Build SVG content
    lines: List[str] = []
    lines.append(f"<svg xmlns='http://www.w3.org/2000/svg' width='{int(W)}' height='{int(H)}' viewBox='0 0 {int(W)} {int(H)}'>")
    lines.append("<style> .grid{stroke:#444;stroke-width:1;fill:none} .hdr{font-weight:bold} .txt{font-family:Arial;font-size:12px;fill:#111} .box{fill:#f7f7f7;stroke:#bbb;stroke-width:1} </style>")
    # Draw cells per row and column
    min_y = min(float(t["y0"]) for t in tokens)
    for r_idx, (y0, y1) in enumerate(row_spans):
        for c_idx in range(len(col_centers)):
            xL = bounds[c_idx]
            xR = bounds[c_idx + 1]
            # Scale to canvas; shift by min_x and min_y to normalize
            x = (xL - min_x) * scale + margin
            y = (y0 - min_y) * scale + margin
            w = max(1.0, (xR - xL) * scale)
            h = max(1.0, (y1 - y0) * scale)
            lines.append(f"<rect class='box' x='{x:.1f}' y='{y:.1f}' width='{w:.1f}' height='{h:.1f}'/>")
    # Draw texts
    for r_idx, cells in enumerate(rows):
        for c_idx, text in enumerate(cells):
            if not text:
                continue
            # Compute text anchor at top-left of the cell box with small padding
            y0, y1 = row_spans[r_idx]
            xL = bounds[c_idx]
            x = (xL - min_x) * scale + margin + 4
            y = (y0 - min_y) * scale + margin + 14
            text_esc = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") )
            lines.append(f"<text class='txt' x='{x:.1f}' y='{y:.1f}'>" + text_esc + "</text>")
    # Draw vertical boundaries as reference
    for b in bounds:
        x = (b - min_x) * scale + margin
        lines.append(f"<line class='grid' x1='{x:.1f}' y1='{margin:.1f}' x2='{x:.1f}' y2='{H - margin:.1f}'/>")
    lines.append("</svg>")
    svg_path = out_prefix.with_suffix(".table.svg")
    svg_path.write_text("\n".join(lines), encoding="utf-8")
    return svg_path


def main() -> None:
    ap = argparse.ArgumentParser(description="Rebuild a table from EasyOCR tokens JSON and export CSV/XLSX/SVG")
    ap.add_argument("--tokens-json", required=True, help="Path to tokens JSON (per-page) from easyocr_dump_spatial.py")
    ap.add_argument("--out-prefix", required=True, help="Output prefix (without extension)")
    ap.add_argument("--row-factor", type=float, default=0.6, help="Row cluster threshold multiplier of median token height")
    ap.add_argument("--gap-mult", type=float, default=2.5, help="Horizontal gap multiplier of median token width to split cell segments")
    ap.add_argument("--center-thresh-mult", type=float, default=2.0, help="Column center clustering threshold multiplier of median token width")
    ap.add_argument("--use-first-row-as-header", action="store_true", help="Use first reconstructed row as header labels")
    ap.add_argument("--svg-scale", type=float, default=0.25, help="Scale for SVG visualization (0.1..1.0)")
    args = ap.parse_args()

    tokens_path = Path(args.tokens_json)
    out_prefix = Path(args.out_prefix)
    tokens = load_tokens(tokens_path)
    if not tokens:
        raise SystemExit("[ERROR] No tokens loaded")

    min_x, max_x, med_w, med_h = compute_stats(tokens)
    rows_tok = cluster_rows(tokens, row_factor=args.row_factor)
    headers, table_rows = build_grid(rows_tok, med_w=med_w, min_x=min_x, max_x=max_x,
                                     gap_mult=args.gap_mult, center_thresh_mult=args.center_thresh_mult)

    # Optional: first row as header
    if args.use_first_row_as_header and table_rows:
        hdr = table_rows[0]
        headers = [h if (h and h.strip()) else f"col_{i+1}" for i, h in enumerate(hdr)]
        table_rows = table_rows[1:]

    csv_path = write_csv(out_prefix, headers, table_rows)
    xlsx_path = write_xlsx(out_prefix, headers, table_rows)
    svg_path = write_svg(out_prefix, headers, table_rows, tokens,
                         row_factor=args.row_factor, gap_mult=args.gap_mult,
                         center_thresh_mult=args.center_thresh_mult, scale=args.svg_scale)

    print(f"[OK] CSV:  {csv_path}")
    if xlsx_path:
        print(f"[OK] XLSX: {xlsx_path}")
    else:
        print("[WARN] openpyxl not installed; skipped XLSX. CSV is available.")
    print(f"[OK] SVG:  {svg_path}")


if __name__ == "__main__":
    main()

