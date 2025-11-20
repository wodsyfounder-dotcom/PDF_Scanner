#!/usr/bin/env python3
"""
Generate a user-friendly Excel schema with dropdowns and conditional formatting:
- Sheet "Template" with headers + validation lists and greying of non-applicable fields by Mode/Return
- Sheet "Instructions" describing each column and examples

Writes: user_inputs/terms.schema.xlsx
"""
from pathlib import Path
import sys


def main() -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
    except Exception:
        print("[ERROR] openpyxl is required to generate the template. Install with: py -m pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    root = Path(__file__).resolve().parents[1]
    out_dir = root / 'user_inputs'
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / 'terms.schema.xlsx'

    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = 'Template'

    # Headers
    headers = [
        'Term', 'Secondary Term', 'Pages', 'Mode',
        'Units', 'Range (min)', 'Range (max)', 'Format', 'GroupAfter', 'GroupBefore',
        'Smart Snap Type', 'Smart Position',
        'OCR_Row_EPS', 'DPI'
    ]
    ws.append(headers)
    header_font = Font(bold=True)
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        ws.column_dimensions[cell.column_letter].width = max(12, len(h) + 2)

    # Data validation lists
    dv_mode = DataValidation(type="list", formula1='"nearest,smart"', allow_blank=True)
    dv_smart = DataValidation(type="list", formula1='"auto,number,date,time,title"', allow_blank=True)
    ws.add_data_validation(dv_mode)
    ws.add_data_validation(dv_smart)
    mode_col_idx = headers.index('Mode') + 1  # Column D (4)
    smart_type_col_idx = headers.index('Smart Snap Type') + 1  # Column K (11)
    secondary_col_idx = headers.index('Secondary Term') + 1  # Column B (2)
    smart_pos_col_idx = headers.index('Smart Position') + 1  # Column L (12)
    units_col_idx = headers.index('Units') + 1  # Column E (5)
    range_min_col_idx = headers.index('Range (min)') + 1  # Column F (6)
    range_max_col_idx = headers.index('Range (max)') + 1  # Column G (7)

    dv_mode.add(f'{chr(64 + mode_col_idx)}2:{chr(64 + mode_col_idx)}2000')
    dv_smart.add(f'{chr(64 + smart_type_col_idx)}2:{chr(64 + smart_type_col_idx)}2000')

    # Conditional formatting to grey non-applicable fields
    grey = PatternFill(start_color='00DDDDDD', end_color='00DDDDDD', fill_type='solid')
    # If Mode <> smart, grey Smart Snap Type (K), Secondary Term (B), Smart Position (L)
    ws.conditional_formatting.add(f'{chr(64 + smart_type_col_idx)}2:{chr(64 + smart_type_col_idx)}2000',
                                   FormulaRule(formula=[f"${chr(64 + mode_col_idx)}2<>\"smart\""], fill=grey))
    ws.conditional_formatting.add(f'{chr(64 + secondary_col_idx)}2:{chr(64 + secondary_col_idx)}2000',
                                   FormulaRule(formula=[f"${chr(64 + mode_col_idx)}2<>\"smart\""], fill=grey))
    ws.conditional_formatting.add(f'{chr(64 + smart_pos_col_idx)}2:{chr(64 + smart_pos_col_idx)}2000',
                                   FormulaRule(formula=[f"${chr(64 + mode_col_idx)}2<>\"smart\""], fill=grey))
    # If Smart Snap Type = "title", grey Units (E), Range (min) (F), Range (max) (G)
    ws.conditional_formatting.add(f'{chr(64 + units_col_idx)}2:{chr(64 + units_col_idx)}2000',
                                   FormulaRule(formula=[f"${chr(64 + smart_type_col_idx)}2=\"title\""], fill=grey))
    ws.conditional_formatting.add(f'{chr(64 + range_min_col_idx)}2:{chr(64 + range_min_col_idx)}2000',
                                   FormulaRule(formula=[f"${chr(64 + smart_type_col_idx)}2=\"title\""], fill=grey))
    ws.conditional_formatting.add(f'{chr(64 + range_max_col_idx)}2:{chr(64 + range_max_col_idx)}2000',
                                   FormulaRule(formula=[f"${chr(64 + smart_type_col_idx)}2=\"title\""], fill=grey))

    # Example rows
    # Column order: Term, Secondary Term, Pages, Mode, Units, Range (min), Range (max), Format, GroupAfter, GroupBefore,
    #               Smart Snap Type, Smart Position, OCR_Row_EPS, DPI
    examples = [
        ['Title', '', '1', 'nearest', '', '', '', 'tpl-xxxx', '', '', '', '', '', ''],
        ['Measured Torque', 'Peak', '3', 'smart', 'lbf', '10', '5000', '', 'Torque Section', '', 'number', '2', '12.0', '850'],
        ['ISP Value', '', '2', 'smart', 'sec', '40', '44', '', 'Pre Test', 'Post Test', 'number', '', '', ''],
    ]
    for row in examples:
        ws.append(row)

    ws.freeze_panes = 'A2'

    # Instructions
    inst = wb.create_sheet('Instructions')
    inst['A1'] = 'EIDP Terms Schema Template'
    inst['A1'].font = Font(bold=True)
    lines = [
        'Fill rows in the Template sheet. Rows appear in outputs in the same order.',
        '',
        'Columns:',
        ' - Term: Friendly row name (required).',
        ' - Pages: Page ranges (1-indexed), e.g., 5-10; 22 (required).',
        ' - Mode: nearest | smart (defaults to nearest).',
        '   * nearest  : Finds the closest numeric value to the term.',
        '   * smart    : Smart Snap mode — finds row by Term and returns the most plausible value to the right.',
        '                - Smart Snap Type: auto | number | date | time | title. Auto-detects if blank.',
        '                - Secondary Term: optional label to validate near the chosen value (sweeps up to page top).',
        '                - Smart Position: optional 1-based index to pick the Nth value to the right (overrides scoring).',
        '',
        'Validation & Filtering:',
        ' - Units: Expected units (e.g., psi, sec, kg) for numeric values.',
        ' - Range (min)/(max): Numeric bounds (scientific notation allowed).',
        ' - Format: Optional value mask for string returns (e.g., tpl-xxxx) or /regex/ if needed.',
        ' - GroupAfter: Only search values after this anchor text appears.',
        ' - GroupBefore: Stop searching once this anchor text is encountered (exclusive).',
        '',
        'OCR Fine-Tuning (optional):',
        ' - OCR_Row_EPS: Y-axis tolerance (0.5-50.0) for grouping OCR text into rows. Default: 8.0.',
        ' - DPI: OCR rendering resolution (e.g., 700, 850). Higher values improve accuracy for small text.',
    ]
    for i, t in enumerate(lines, start=2):
        inst[f'A{i}'] = t
        inst[f'A{i}'].alignment = Alignment(wrap_text=True)
    inst.column_dimensions['A'].width = 110

    wb.save(out_path)
    print(f"[DONE] Wrote template -> {out_path}")


if __name__ == '__main__':
    main()
