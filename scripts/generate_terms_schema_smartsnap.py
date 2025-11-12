#!/usr/bin/env python3
"""
Generate the Smart-Snap terms schema workbook with helpful validation, formatting,
and instructions. Output: user_inputs/terms.schema.smartsnap.xlsx
"""

from __future__ import annotations

from pathlib import Path
import sys


HEADERS = [
    "Data Group",
    "Term Label",
    "Smart Snap Type",
    "Term",
    "Pages",
    "GroupAfter",
    "GroupBefore",
    "Units",
    "Range (min)",
    "Range (max)",
    "Format",
    "Secondary Term",
    "Smart Position",
    # Hidden legacy columns (preserved for backward compatibility)
    "Mode",
    "Line",
    "Column",
    "Anchor",
    "FieldIndex",
    "FieldSplit",
    "Return",
]

# Columns to hide in Excel (legacy full-table mode columns)
HIDDEN_COLUMNS = ["Mode", "Line", "Column", "Anchor", "FieldIndex", "FieldSplit", "Return"]

MODE_OPTIONS = ["smart", "full table"]
RETURN_OPTIONS = ["number", "string"]
SPLIT_OPTIONS = ["auto", "groups", "tokens"]
SMART_TYPES = ["", "auto", "number", "date", "time", "title"]


def main() -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
    except Exception:
        print(
            "[ERROR] openpyxl is required to generate the Smart-Snap schema. "
            "Install with: py -m pip install openpyxl",
            file=sys.stderr,
        )
        sys.exit(1)

    root = Path(__file__).resolve().parents[1]
    out_dir = root / "user_inputs"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "terms.schema.smartsnap.xlsx"

    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = "Template"

    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True)
    ws.append(HEADERS)
    for idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = header_font
        col_letter = get_column_letter(idx)

        # Hide legacy columns
        if header in HIDDEN_COLUMNS:
            ws.column_dimensions[col_letter].hidden = True
            ws.column_dimensions[col_letter].width = 14
        else:
            ws.column_dimensions[col_letter].width = max(14, len(header) + 2)

    # Data validations
    def _dv(values: list[str]) -> DataValidation:
        return DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)

    # Find column indices for validation
    smart_type_col = HEADERS.index("Smart Snap Type") + 1
    mode_col = HEADERS.index("Mode") + 1
    return_col = HEADERS.index("Return") + 1
    split_col = HEADERS.index("FieldSplit") + 1

    dv_mode = _dv(MODE_OPTIONS)
    dv_return = _dv(RETURN_OPTIONS)
    dv_split = _dv(SPLIT_OPTIONS)
    dv_smart = _dv([v or " " for v in SMART_TYPES])

    for dv in (dv_mode, dv_return, dv_split, dv_smart):
        ws.add_data_validation(dv)
    dv_smart.add(f"{get_column_letter(smart_type_col)}2:{get_column_letter(smart_type_col)}2000")
    dv_mode.add(f"{get_column_letter(mode_col)}2:{get_column_letter(mode_col)}2000")
    dv_return.add(f"{get_column_letter(return_col)}2:{get_column_letter(return_col)}2000")
    dv_split.add(f"{get_column_letter(split_col)}2:{get_column_letter(split_col)}2000")

    # No conditional formatting - no grey cells

    # Example rows (Smart Snap mode)
    # Column order: Data Group, Term Label, Smart Snap Type, Term, Pages, GroupAfter, GroupBefore,
    #               Units, Range (min), Range (max), Format, Secondary Term, Smart Position,
    #               Mode, Line, Column, Anchor, FieldIndex, FieldSplit, Return
    examples = [
        ["Metadata", "Title", "title", "title", "1", "", "", "", "", "", "", "", "", "smart", "", "", "", "", "", "string"],
        ["Performance", "Pre-Test ISP", "number", "isp", "1", "Pre Test", "Post Test", "sec", "40", "44", "", "", "", "smart", "", "", "", "", "", "number"],
        ["Tables", "Proof Load", "number", "proof load", "2", "", "", "psi", "", "", "", "", "", "smart", "", "", "", "", "", "number"],
    ]
    for row in examples:
        ws.append(row)

    ws.freeze_panes = "A2"

    inst = wb.create_sheet("Instructions")
    inst["A1"] = "Smart-Snap Terms Schema"
    inst["A1"].font = Font(bold=True, size=13)
    instructions = [
        "Fill rows in the Template sheet. Rows are processed in order using Smart-Snap extraction.",
        "",
        "Required columns:",
        " - Term: The search term to find in the PDF (e.g., 'title', 'proof load', 'isp').",
        " - Pages: One or more ranges (1-indexed). Examples: '1', '1-3', '5-7; 10; 12-13'.",
        "",
        "Data Organization:",
        " - Data Group: Organize extracted data into categories (e.g., 'Metadata', 'Performance', 'Tables').",
        " - Term Label: Friendly label for the extracted value in reports.",
        "",
        "Search Configuration:",
        " - Smart Snap Type: Hint for extraction type - 'number', 'date', 'time', 'title', or blank for auto-detect.",
        " - GroupAfter / GroupBefore: Limit search to text appearing after/before these anchor terms.",
        "",
        "Validation & Scoring:",
        " - Units: Expected units (e.g., 'psi', 'sec', 'kg').",
        " - Range (min) / Range (max): Valid range for numeric values.",
        " - Format: Optional pattern or /regex/ for validation.",
        "",
        "Advanced Extraction:",
        " - Secondary Term: Disambiguate values when multiple matches exist in a row.",
        " - Smart Position: Position hint for value extraction.",
        "",
        "Tips:",
        " - Leave columns blank when not needed.",
        " - Use consistent Data Groups and Term Labels to organize your dashboards.",
        " - Save as .xlsx; the scanner also accepts .csv with the same column headers.",
    ]
    for idx, line in enumerate(instructions, start=3):
        cell = inst[f"A{idx}"]
        cell.value = line
        cell.alignment = Alignment(wrap_text=True)
    inst.column_dimensions["A"].width = 110

    wb.save(out_path)
    print(f"[DONE] Smart-Snap schema written -> {out_path}")


if __name__ == "__main__":
    main()
